"""Update member details from 2026春季全息表 Excel.

Reads the combined holistic table and updates existing Member records.
Reports all changes in detail. Creates new members if not found in DB.

Usage:
    cd backend && python update_member_from_holistic.py [--dry-run]
"""
import sys
import os
from datetime import datetime, date

import re
import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from app import create_app
from database import db
from models import Member, User

EXCEL_PATH = os.path.join(
    os.path.dirname(os.path.abspath(__file__)),
    '..', 'assets',
    '2026年春季学期清华大学学生艺术团舞蹈队-全息表统计(2).xlsx',
)

# Column mapping: 0-indexed column → field name
COL_MAP = {
    0: 'name',
    1: 'gender',
    2: 'student_id',
    3: 'department',
    4: 'class_name',
    5: 'phone',
    6: 'email',
    7: 'dormitory',
    8: 'birth_date',
    9: 'ethnicity',
    10: 'hometown',
    11: 'political_status',
    12: 'party_branch',
    13: 'is_talented',
    14: 'is_concentrated_class',
    15: 'team_role',
    16: 'join_year',
    17: 'team_level',
    18: 'graduating_this_semester',
}

BOOL_FIELDS = {'is_talented', 'is_concentrated_class', 'graduating_this_semester'}
DATE_FIELDS = {'birth_date'}
INT_FIELDS = {'join_year'}
SKIP_FIELDS = {'name'}  # name is used for matching, not updating


def parse_bool(val):
    if val is None:
        return False
    s = str(val).strip()
    return s in ('是', 'True', 'true', '1', 'yes')


def parse_date(val):
    if val is None:
        return None
    if isinstance(val, (datetime, date)):
        return val if isinstance(val, date) and not isinstance(val, datetime) else val.date() if isinstance(val, datetime) else val
    s = str(val).strip()
    for fmt in ('%Y-%m-%d', '%Y/%m/%d', '%Y.%m.%d'):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def parse_int(val):
    if val is None:
        return None
    s = str(val).strip()
    # Strip trailing non-numeric chars (e.g. '2022年' → '2022')
    s = re.sub(r'[^\d].*$', '', s)
    if not s:
        return None
    try:
        v = int(s)
        # Fix 2-digit years (e.g. 21 → 2021)
        if v < 100:
            v += 2000
        return v
    except (ValueError, TypeError):
        return None


def parse_string(val):
    if val is None:
        return None
    s = str(val).strip()
    return s if s else None


def normalize_team_role(val):
    """Normalize team_role: '无', '否', '新入队.暂无' → None."""
    if val is None:
        return None
    s = str(val).strip()
    if s in ('无', '否', '新入队.暂无', ''):
        return None
    return s


def main():
    dry_run = '--dry-run' in sys.argv

    if not os.path.exists(EXCEL_PATH):
        print(f'错误: 文件不存在: {EXCEL_PATH}')
        return

    app = create_app()

    with app.app_context():
        member_map = {m.name: m for m in Member.query.all()}
        db_names = set(member_map.keys())

        wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
        ws = wb['Sheet1']

        excel_names = set()
        seen_names = set()
        updated_count = 0
        created_count = 0
        skipped_duplicates = []
        warnings = []
        all_changes = []

        for row_idx in range(2, ws.max_row + 1):
            # Read row data
            row_data = {}
            for col_idx, field in COL_MAP.items():
                raw_val = ws.cell(row=row_idx, column=col_idx + 1).value
                row_data[field] = raw_val

            name = parse_string(row_data.get('name'))
            if not name:
                continue

            # Handle duplicates in Excel
            if name in seen_names:
                skipped_duplicates.append(f'  跳过重复行: {name} (行 {row_idx})')
                continue
            seen_names.add(name)
            excel_names.add(name)

            # Parse all fields
            parsed = {}
            for field, raw_val in row_data.items():
                if field in SKIP_FIELDS:
                    continue
                if field in BOOL_FIELDS:
                    parsed[field] = parse_bool(raw_val)
                elif field in DATE_FIELDS:
                    parsed[field] = parse_date(raw_val)
                elif field in INT_FIELDS:
                    parsed[field] = parse_int(raw_val)
                elif field == 'team_role':
                    parsed[field] = normalize_team_role(raw_val)
                elif field == 'phone':
                    # Phone numbers may come as floats from Excel
                    if raw_val is not None:
                        parsed[field] = str(int(raw_val)) if isinstance(raw_val, (int, float)) else str(raw_val).strip()
                    else:
                        parsed[field] = None
                else:
                    parsed[field] = parse_string(raw_val)

            # Fix known birth_date errors in Excel
            BIRTH_DATE_FIXES = {
                '乔炫嘉': date(2005, 9, 9),
                '王雪莹': date(2006, 2, 9),
            }
            if name in BIRTH_DATE_FIXES:
                parsed['birth_date'] = BIRTH_DATE_FIXES[name]

            # Data warnings
            if parsed.get('birth_date'):
                bd = parsed['birth_date']
                if bd.year > 2010:
                    warnings.append(f'  ⚠ {name}: birth_date={bd} 看起来不对 (年份>2010)')
            if parsed.get('join_year') and parsed['join_year'] > 2026:
                warnings.append(f'  ⚠ {name}: join_year={parsed["join_year"]} 看起来不对')

            member = member_map.get(name)

            if member:
                # Update existing member
                changes = []
                for field, new_val in parsed.items():
                    old_val = getattr(member, field, None)
                    if old_val != new_val:
                        changes.append(f'    {field}: {repr(old_val)} → {repr(new_val)}')
                        if not dry_run:
                            setattr(member, field, new_val)

                if changes:
                    all_changes.append(f'  更新 {name}:')
                    all_changes.extend(changes)
                    updated_count += 1
            else:
                # Create new member
                all_changes.append(f'  新建 {name} (Excel 中有, DB 中无)')
                if not dry_run:
                    new_member = Member(name=name, status='active')
                    for field, val in parsed.items():
                        setattr(new_member, field, val)
                    db.session.add(new_member)
                    db.session.flush()
                    # Create user account
                    if new_member.student_id:
                        existing_user = User.query.filter_by(username=name).first()
                        if not existing_user:
                            sid = new_member.student_id.strip()
                            password = sid[-6:] if len(sid) >= 6 else sid
                            user = User(
                                username=name,
                                display_name=name,
                                role=User.ROLE_MEMBER,
                                status='active',
                                member_id=new_member.id,
                                email=new_member.email,
                                phone=new_member.phone,
                            )
                            user.set_password(password)
                            db.session.add(user)
                            all_changes.append(f'    → 创建用户账号 (密码: 学号后6位)')
                created_count += 1

        wb.close()

        # Members in DB but not in Excel
        missing_from_excel = db_names - excel_names
        active_missing = [n for n in missing_from_excel if member_map[n].status == 'active']

        if not dry_run:
            db.session.commit()

        # Report
        mode = '[DRY RUN] ' if dry_run else ''
        print(f'\n{mode}=== 全息表数据更新报告 ===')
        print(f'Excel 成员数: {len(excel_names)} (去重后)')
        print(f'数据库成员数: {len(db_names)}')
        print(f'更新: {updated_count} 位 | 新建: {created_count} 位')

        if skipped_duplicates:
            print(f'\n--- 重复行 ({len(skipped_duplicates)}) ---')
            for line in skipped_duplicates:
                print(line)

        if warnings:
            print(f'\n--- 数据警告 ({len(warnings)}) ---')
            for line in warnings:
                print(line)

        if all_changes:
            print(f'\n--- 变更明细 ---')
            for line in all_changes:
                print(line)

        if active_missing:
            print(f'\n--- DB 中在队但不在 Excel 中 ({len(active_missing)}) ---')
            for name in sorted(active_missing):
                m = member_map[name]
                print(f'  {name} (id={m.id}, 梯队={m.team_level}, 入队={m.join_year})')

        print(f'\n{mode}完成。')


if __name__ == '__main__':
    main()
